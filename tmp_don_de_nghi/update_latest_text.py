from pathlib import Path
from copy import deepcopy
from shutil import copy2
from zipfile import ZipFile
import sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
target = root / 'Don_de_nghi_giai_quyet_ho_so_chuyen_muc_dich_su_dung_dat.docx'
backup = Path(__file__).parent / 'before_latest_text.docx'
if not backup.exists():
    copy2(target, backup)
doc = Document(target)
body = doc._element.body
for item in list(body):
    if item.tag != qn('w:sectPr'):
        body.remove(item)

expected = []
def add(text, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False,
        size=14, indent=True, after=6, keep=False, pagebreak=False,
        left=0, style=None, highlights=()):
    expected.append(text)
    p = doc.add_paragraph(style=style)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.first_line_indent = Cm(1 if indent else 0)
    pf.left_indent = Cm(left)
    pf.keep_together = True
    pf.keep_with_next = keep
    pf.page_break_before = pagebreak
    pf.widow_control = True
    remaining = text
    chunks = []
    for fragment in highlights:
        i = remaining.index(fragment)
        chunks.append((remaining[:i], bold))
        chunks.append((fragment, True))
        remaining = remaining[i + len(fragment):]
    chunks.append((remaining, bold))
    for content, is_bold in chunks:
        if not content:
            continue
        r = p.add_run(content)
        r.bold, r.italic = is_bold, italic
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0, 0, 0)
        rpr = r._element.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + key), 'Times New Roman')
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), 'vi-VN')
        rpr.append(lang)
    return p

C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT
add('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', align=C, bold=True, size=13, indent=False, after=0, keep=True)
add('Độc lập - Tự do - Hạnh phúc', align=C, bold=True, indent=False, after=13, keep=True)
add('ĐƠN ĐỀ NGHỊ', align=C, bold=True, size=16, indent=False, after=3, keep=True, style='Title')
add('Về việc xem xét và thông tin tình hình giải quyết\nhồ sơ chuyển mục đích sử dụng đất', align=C, bold=True, indent=False, after=12, keep=True)
add('Kính gửi: Chủ tịch Ủy ban nhân dân phường Duy Tân, tỉnh Ninh Bình', align=C, bold=True, indent=False, after=10, keep=True)

for text in [
    'Tôi tên là: .........................................................................................................',
    'Ngày sinh: ....../....../............\u2003\u2003Số CCCD: ....................................................',
    'Địa chỉ thường trú: ...............................................................................................',
    'Số điện thoại: ......................................................................................................',
    'Thửa đất số: ................................., tờ bản đồ số: ..................................................',
    'Địa chỉ thửa đất: ...................................................................................................',
]:
    add(text, align=L, indent=False, after=5)

add('Gia đình tôi hiện đang sử dụng thửa đất nêu trên có tổng diện tích 655 m², trong đó có 106 m² đất trồng cây lâu năm (CLN). Gia đình tôi có nguyện vọng chuyển phần diện tích 106 m² này sang đất ở tại đô thị (ODT) để phục vụ nhu cầu ổn định chỗ ở của gia đình.', highlights=['655 m²', '106 m² đất trồng cây lâu năm (CLN)', 'đất ở tại đô thị (ODT)'])
add('Từ đầu tháng 05/2026, tôi đã nhiều lần liên hệ và làm việc với bộ phận chuyên môn để tìm hiểu, được hướng dẫn và thực hiện thủ tục chuyển mục đích sử dụng đất.')
add('Ngày 06/10/2026, tôi đã nộp hồ sơ đề nghị chuyển mục đích sử dụng đất và được cán bộ chuyên môn tiếp nhận để xử lý. Sau thời điểm này, tôi tiếp tục chủ động liên hệ để nắm tình hình giải quyết hồ sơ.', highlights=['06/10/2026'])
add('Tuy nhiên, đến nay tôi chưa nhận được thông tin cụ thể về hồ sơ đang được xử lý đến bước nào, có nội dung nào cần xác minh, bổ sung hoặc còn khó khăn, vướng mắc cần giải quyết hay không.')
add('Quá trình thực hiện thủ tục đã kéo dài trong một khoảng thời gian đáng kể. Gia đình tôi đã dành nhiều thời gian, công sức để liên hệ, đi lại và theo dõi nhưng đến nay vẫn chưa xác định được hướng xử lý và thời điểm dự kiến hoàn thành. Điều này ảnh hưởng trực tiếp đến kế hoạch ổn định chỗ ở của gia đình tôi.')
add('Trên cơ sở Quyết định số 64/2026/QĐ-UBND ngày 10/6/2026 của Ủy ban nhân dân tỉnh Ninh Bình quy định về trình tự, thủ tục hành chính lĩnh vực đất đai trên địa bàn tỉnh và các quy định pháp luật có liên quan, tôi kính mong Chủ tịch Ủy ban nhân dân phường Duy Tân quan tâm, xem xét tình hình giải quyết hồ sơ nêu trên.', pagebreak=True, highlights=['Quyết định số 64/2026/QĐ-UBND ngày 10/6/2026 của Ủy ban nhân dân tỉnh Ninh Bình', 'Chủ tịch Ủy ban nhân dân phường Duy Tân quan tâm, xem xét tình hình giải quyết hồ sơ nêu trên'])
add('Tôi mong được thông tin bằng văn bản về tình trạng và hướng giải quyết hồ sơ để gia đình có cơ sở theo dõi và chủ động thực hiện các công việc có liên quan.')
add('Trường hợp hồ sơ đủ điều kiện tiếp tục giải quyết, tôi mong được biết các bước còn lại và thời gian dự kiến hoàn thành. Trường hợp hồ sơ chưa đủ điều kiện hoặc không đủ điều kiện giải quyết, tôi mong được biết lý do, căn cứ pháp lý và nội dung cần thực hiện thêm (nếu có).')
add('Gia đình tôi sẵn sàng phối hợp, cung cấp và bổ sung đầy đủ thông tin, tài liệu theo yêu cầu trong quá trình giải quyết hồ sơ.', keep=True)
add('Tôi xin trân trọng cảm ơn!', bold=True, after=10, keep=True)
add('Duy Tân, ngày ...... tháng ...... năm 2026', align=C, italic=True, size=13, indent=False, left=7, after=3, keep=True)
add('NGƯỜI LÀM ĐƠN', align=C, bold=True, size=13, indent=False, left=7, after=1, keep=True)
add('(Ký và ghi rõ họ tên)', align=C, italic=True, size=13, indent=False, left=7, after=46, keep=True)
add('...........................................', align=C, indent=False, left=7, after=0)
doc.core_properties.subject = 'Đề nghị xem xét và thông tin tình hình giải quyết hồ sơ chuyển mục đích sử dụng đất'
doc.core_properties.last_modified_by = ''
doc.save(target)

verified = Document(target)
assert [p.text for p in verified.paragraphs] == expected
assert len(expected) == 25
with ZipFile(target) as z:
    assert z.testzip() is None
    assert all(b'Aspose' not in z.read(n) for n in z.namelist() if n.endswith('.xml'))
print('Verified full supplied text, 25 paragraphs, and DOCX integrity.')

sys.path.insert(0, 'C:/Users/thuan/.codex/visualizations/2026/10/19/01a153ab-84a4-71c3-a2f0-9c7a76c0f6b8/render_dependencies')
import aspose.words as aw
preview = aw.Document(str(target))
print('Page count:', preview.page_count)
out = Path(__file__).parent / 'latest_render'
out.mkdir(exist_ok=True)
preview.save(str(out / 'preview.pdf'))
