from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'Don_de_nghi_giai_quyet_ho_so_chuyen_muc_dich_su_dung_dat.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.header_distance = sec.footer_distance = Cm(0.8)

for name in ['Normal', 'Title']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), 'Times New Roman')
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        for border in list(ppr.findall(qn('w:pBdr'))):
            ppr.remove(border)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.15
normal.space_after = Pt(6)
normal.widow_control = True

def para(text='', *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False,
         size=14, before=0, after=6, indent=True, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.first_line_indent = Cm(1) if indent else Cm(0)
    f.keep_together = True
    f.keep_with_next = keep
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    r.font.name = 'Times New Roman'
    r.font.color.rgb = RGBColor(0,0,0)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'vi-VN')
    r._element.get_or_add_rPr().append(lang)
    return p

C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT
para('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', align=C, bold=True, size=13, after=0, indent=False, keep=True)
para('Độc lập - Tự do - Hạnh phúc', align=C, bold=True, after=13, indent=False, keep=True)
para('ĐƠN ĐỀ NGHỊ', align=C, bold=True, size=16, after=3, indent=False, keep=True, style='Title')
para('Về việc quan tâm, xem xét và chỉ đạo giải quyết\nhồ sơ chuyển mục đích sử dụng đất', align=C, bold=True, size=14, after=12, indent=False, keep=True)
para('Kính gửi: Chủ tịch Ủy ban nhân dân phường Duy Tân, tỉnh Ninh Bình', align=C, bold=True, after=10, indent=False, keep=True)

fields = [
    'Tôi tên là: ' + '.' * 105,
    'Ngày sinh: ....../....../............      Số CCCD: ' + '.' * 32,
    'Địa chỉ thường trú: ' + '.' * 95,
    'Số điện thoại: ' + '.' * 102,
    'Thửa đất số: ................................., tờ bản đồ số: ' + '.' * 50,
    'Địa chỉ thửa đất: ' + '.' * 99,
    'Mã hồ sơ/Số phiếu tiếp nhận (nếu có): ' + '.' * 50,
    'Ngày hẹn trả kết quả (nếu có): ....../....../............',
]
for s in fields:
    para(s, align=L, after=5, indent=False)

para('Gia đình tôi hiện đang sử dụng thửa đất nêu trên có tổng diện tích 655 m², trong đó có 106 m² đất trồng cây lâu năm (CLN). Gia đình tôi có nhu cầu chuyển mục đích sử dụng 106 m² đất CLN sang đất ở tại đô thị (ODT).')
para('Ngày 06/10/2026, tôi đã nộp hồ sơ đề nghị chuyển mục đích sử dụng đất và được cán bộ chuyên môn tiếp nhận để giải quyết. Trong thời gian chờ giải quyết, tôi đã chủ động liên hệ, trực tiếp trao đổi với cán bộ chuyên môn và lãnh đạo bộ phận có liên quan để nắm tình hình, phối hợp xử lý hồ sơ.')
para('Tôi hiểu và chia sẻ với khối lượng công việc của địa phương, cũng như việc giải quyết hồ sơ đất đai có thể cần nhiều bước và sự phối hợp giữa các cơ quan, đơn vị. Tuy nhiên, đến nay tôi chưa được biết cụ thể hồ sơ đang được giải quyết đến bước nào, có nội dung cần xác minh hoặc bổ sung hay không. Việc giải quyết hồ sơ có ý nghĩa quan trọng đối với kế hoạch ổn định chỗ ở của gia đình tôi.')
p = para('Trên cơ sở Quyết định số 64/2026/QĐ-UBND ngày 10/6/2026 của Ủy ban nhân dân tỉnh Ninh Bình quy định về trình tự, thủ tục hành chính lĩnh vực đất đai trên địa bàn tỉnh và các quy định pháp luật có liên quan, tôi kính đề nghị Chủ tịch Ủy ban nhân dân phường Duy Tân quan tâm xem xét, chỉ đạo:', keep=True)
p.paragraph_format.page_break_before = True

requests = [
    '1. Xem xét, nắm bắt tình hình và tiến độ giải quyết hồ sơ chuyển mục đích sử dụng đất của gia đình tôi đã nộp ngày 06/10/2026.',
    '2. Chỉ đạo bộ phận chuyên môn tiếp tục thực hiện các bước cần thiết để hồ sơ được giải quyết theo đúng trình tự, thủ tục và thời hạn quy định; thông tin cho tôi về tiến độ và thời gian dự kiến trả kết quả.',
    '3. Trường hợp hồ sơ cần xác minh, bổ sung hoặc có khó khăn, vướng mắc, kính đề nghị thông tin rõ nội dung và hướng dẫn để gia đình tôi kịp thời phối hợp thực hiện.',
]
for s in requests:
    p = para(s, indent=False)
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)

para('Gia đình tôi xin nghiêm túc phối hợp, cung cấp thông tin, tài liệu và thực hiện đầy đủ các yêu cầu theo quy định trong quá trình giải quyết hồ sơ.', keep=True)
para('Tôi xin trân trọng cảm ơn!', bold=True, after=10, keep=True)
for text, bold, italic, after in [
    ('Duy Tân, ngày ...... tháng ...... năm 2026', False, True, 3),
    ('NGƯỜI LÀM ĐƠN', True, False, 1),
    ('(Ký và ghi rõ họ tên)', False, True, 46),
]:
    p = para(text, align=C, bold=bold, italic=italic, size=13, after=after, indent=False, keep=True)
    p.paragraph_format.left_indent = Cm(7)
p = para('.' * 43, align=C, indent=False, after=0)
p.paragraph_format.left_indent = Cm(7)

doc.core_properties.title = 'Đơn đề nghị giải quyết hồ sơ chuyển mục đích sử dụng đất'
doc.core_properties.subject = 'Đề nghị quan tâm, xem xét tiến độ và chỉ đạo giải quyết hồ sơ'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.comments = ''
doc.save(OUT)
print(OUT)
