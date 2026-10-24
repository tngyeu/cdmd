from pathlib import Path
from zipfile import ZipFile
import sys
from docx import Document

root = Path(__file__).resolve().parent.parent
path = root / 'Don_de_nghi_giai_quyet_ho_so_chuyen_muc_dich_su_dung_dat.docx'
doc = Document(path)
before = [p.text for p in doc.paragraphs]
replacements = {
    'Ngày 06/10/2026, tôi đã nộp hồ sơ đề nghị chuyển mục đích sử dụng đất và được cán bộ chuyên môn tiếp nhận để giải quyết. Trong thời gian chờ giải quyết, tôi đã chủ động liên hệ, trực tiếp trao đổi với cán bộ chuyên môn và lãnh đạo bộ phận có liên quan để nắm tình hình, phối hợp xử lý hồ sơ.':
    'Ngày 06/10/2026, tôi đã nộp hồ sơ đề nghị chuyển mục đích sử dụng đất và được cán bộ chuyên môn tiếp nhận để giải quyết. Trong thời gian chờ giải quyết, tôi đã chủ động liên hệ, trao đổi với bộ phận chuyên môn để nắm tình hình và phối hợp trong quá trình xử lý hồ sơ.',
    'Tôi hiểu và chia sẻ với khối lượng công việc của địa phương, cũng như việc giải quyết hồ sơ đất đai có thể cần nhiều bước và sự phối hợp giữa các cơ quan, đơn vị. Tuy nhiên, đến nay tôi chưa được biết cụ thể hồ sơ đang được giải quyết đến bước nào, có nội dung cần xác minh hoặc bổ sung hay không. Việc giải quyết hồ sơ có ý nghĩa quan trọng đối với kế hoạch ổn định chỗ ở của gia đình tôi.':
    'Tôi hiểu hiện nay khối lượng công việc của địa phương tương đối lớn, đồng thời việc giải quyết hồ sơ đất đai có thể cần thực hiện nhiều bước và có sự phối hợp giữa các cơ quan, đơn vị. Tuy nhiên, đến nay tôi chưa được biết cụ thể hồ sơ đang được giải quyết đến bước nào, có nội dung cần xác minh, bổ sung hoặc phối hợp thực hiện hay không. Việc giải quyết hồ sơ có ý nghĩa quan trọng đối với kế hoạch ổn định chỗ ở của gia đình tôi.',
    '2. Chỉ đạo bộ phận chuyên môn tiếp tục thực hiện các bước cần thiết để hồ sơ được giải quyết theo đúng trình tự, thủ tục và thời hạn quy định; thông tin cho tôi về tiến độ và thời gian dự kiến trả kết quả.':
    '2. Chỉ đạo bộ phận chuyên môn tiếp tục thực hiện các bước cần thiết để hồ sơ được giải quyết theo đúng trình tự, thủ tục và thời hạn quy định; đồng thời thông tin để tôi được biết về tình trạng xử lý hồ sơ và dự kiến bước thực hiện tiếp theo.',
}
for old in replacements:
    assert before.count(old) == 1, 'Expected exactly one matching paragraph'
for p in doc.paragraphs:
    if p.text in replacements:
        new = replacements[p.text]
        assert len(p.runs) == 1, 'Preserve the existing single-run formatting'
        p.runs[0].text = new
doc.save(path)
verified = Document(path)
after = [p.text for p in verified.paragraphs]
assert len(before) == len(after)
assert sum(a != b for a, b in zip(before, after)) == 3
assert all(b == replacements.get(a, a) for a, b in zip(before, after))
assert any('Mã hồ sơ/Số phiếu tiếp nhận (nếu có)' in p for p in after)
assert any('Ngày hẹn trả kết quả (nếu có)' in p for p in after)
with ZipFile(path) as z:
    assert z.testzip() is None
    assert all(b'Aspose' not in z.read(n) for n in z.namelist() if n.endswith('.xml'))
print('Verified exactly three paragraph edits; other content and formatting preserved.')

sys.path.insert(0, 'C:/Users/thuan/.codex/visualizations/2026/10/19/01a153ab-84a4-71c3-a2f0-9c7a76c0f6b8/render_dependencies')
import aspose.words as aw
render = root / 'tmp_don_de_nghi' / 'revised_render'
render.mkdir(exist_ok=True)
preview = aw.Document(str(path))
print('Rendered pages:', preview.page_count)
assert preview.page_count == 2
preview.save(str(render / 'preview.pdf'))
